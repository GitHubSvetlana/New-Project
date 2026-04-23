import openpyxl
import os

if os.path.exists('names.xlsx'):
    workbook = openpyxl.load_workbook('names.xlsx')
    sheet = workbook.active
    if sheet['C1'].value is None:
        sheet['C1'] = 'Ticket Number'
    next_row = sheet.max_row + 1
else:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet['A1'] = 'First Name'
    sheet['B1'] = 'Last Name'
    sheet['C1'] = 'Ticket Number'
    next_row = 2

first_name = input("Enter your first name: ")
last_name = input("Enter your last name: ")
ticket_number = input("Enter your ticket number: ")

sheet[f'A{next_row}'] = first_name
sheet[f'B{next_row}'] = last_name
sheet[f'C{next_row}'] = ticket_number

workbook.save('names.xlsx')

print("Excel file 'names.xlsx' has been updated with your details!")

os.startfile('names.xlsx')
